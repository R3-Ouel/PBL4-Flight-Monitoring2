from fastapi import FastAPI, WebSocket, WebSocketDisconnect, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
import asyncio
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
import shutil
import tempfile
import os

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ConnectionManager:
    def __init__(self):
        self.active_connections: list[WebSocket] = []

    async def connect(self, websocket: WebSocket):
        await websocket.accept()
        self.active_connections.append(websocket)

    def disconnect(self, websocket: WebSocket):
        try:
            self.active_connections.remove(websocket)
        except ValueError:
            pass

    async def broadcast(self, message: dict):
        data = json.dumps(message)
        to_remove = []
        for connection in list(self.active_connections):
            try:
                await connection.send_text(data)
            except Exception:
                to_remove.append(connection)
        for c in to_remove:
            self.disconnect(c)

ws_manager = ConnectionManager()

# Import the CSV/Supabase manager from the local `core` package. Use a different
# name to avoid clobbering the websocket `ws_manager` above and to ensure the
# import works when running from the `backend/` folder (importing `backend.*`
# would fail when the current working directory is already `backend`).
import core.manager as csv_manager


@app.on_event("startup")
async def startup_event():
    try:
        csv_manager.start_csv_pusher(interval=5)
        print("CSV pusher started (interval=5s)")
    except Exception as e:
        print("Failed to start CSV pusher:", e)


@app.on_event("shutdown")
async def shutdown_event():
    try:
        csv_manager.stop_csv_pusher()
    except Exception:
        pass

@app.get("/download-excel")
async def download_excel():
    # Read the CSV data
    csv_path = "data/flight_data.csv"
    df = pd.read_csv(csv_path)
    
    # Create workbook
    wb = Workbook()
    
    # Data sheet
    ws_data = wb.active
    ws_data.title = "Data"
    
    # Write headers (row 1)
    for c, col in enumerate(df.columns, 1):
        ws_data.cell(row=1, column=c, value=col)

    # Write data starting at row 2
    for r, row in enumerate(df.values.tolist(), start=2):
        for c, val in enumerate(row, start=1):
            ws_data.cell(row=r, column=c, value=val)
    
    # Create chart sheets
    charts = [
        ("Altitude", "altitude", 2),
        ("Vitesse", "vitesse", 3),
        ("AZ", "az", 6),
        ("Orientation", "roll,pitch,yaw", [7,8,9]),
    ]
    
    for title, cols, col_indices in charts:
        ws_chart = wb.create_sheet(title=title)
        chart = LineChart()
        chart.title = title
        chart.x_axis.title = "Timestamp"
        chart.y_axis.title = "Value"
        
        if isinstance(col_indices, list):
            for idx in col_indices:
                # include header row in the data reference so the header is used
                # as the series title in the chart
                data = Reference(ws_data, min_col=idx, min_row=1, max_row=len(df)+1)
                cats = Reference(ws_data, min_col=1, min_row=2, max_row=len(df)+1)
                chart.add_data(data, titles_from_data=True)
        else:
            data = Reference(ws_data, min_col=col_indices, min_row=1, max_row=len(df)+1)
            cats = Reference(ws_data, min_col=1, min_row=2, max_row=len(df)+1)
            chart.add_data(data, titles_from_data=True)
        
        chart.set_categories(cats)
        ws_chart.add_chart(chart, "A1")
    
    # Save to temp file
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        tmp_path = tmp.name

    # Try to copy the generated file to the user's Downloads folder (Windows: C:\\Users\\<user>\\Downloads)
    try:
        downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
        os.makedirs(downloads_dir, exist_ok=True)
        dest_path = os.path.join(downloads_dir, "flight_data.xlsx")
        shutil.copy(tmp_path, dest_path)
        # serve the copied file so it's also present in Downloads
        return FileResponse(dest_path, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename='flight_data.xlsx')
    except Exception as e:
        # Fallback to serving the temp file if copying failed
        print("Failed to copy to Downloads:", e)
        return FileResponse(tmp_path, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename='flight_data.xlsx')

@app.websocket("/ws/flight-data")
async def websocket_endpoint(websocket: WebSocket):
    await ws_manager.connect(websocket)
    try:
        while True:
            await websocket.receive_text()  # pour garder la connexion
    except WebSocketDisconnect:
        ws_manager.disconnect(websocket)

@app.post("/push")
async def push(request: Request):
    try:
        payload = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "error": "invalid json"}, status_code=400)

    # Broadcast to websocket clients
    asyncio.create_task(ws_manager.broadcast(payload))

    # Schedule ingestion (CSV + Supabase) in background thread to avoid blocking
    try:
        # Write payload to CSV in a thread to avoid blocking the event loop
        asyncio.create_task(asyncio.to_thread(csv_manager.process_payload, payload))
    except Exception as e:
        print("Failed to schedule ingestion:", e)

    return {"ok": True}